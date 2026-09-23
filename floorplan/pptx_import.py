"""Übernimmt eine PowerPoint-Skizze der Halle als Hallenplan.

Jede Form (Rechteck, Ellipse, …) innerhalb der Folie wird zu einem Planobjekt. Textfelder, Linien
und alles außerhalb der Folie (Notizen, Ablageflächen) werden ignoriert. Standardmaßstab: 1 cm = 1 m.
"""

from lxml import etree
from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from .models import PlanElement

EMU_PER_CM = 360000
Kind = PlanElement.Kind

# Theme-Farbnamen von python-pptx -> Einträge im Farbschema der Datei
THEME_SLOTS = {
    "BACKGROUND_1": "lt1",
    "TEXT_1": "dk1",
    "BACKGROUND_2": "lt2",
    "TEXT_2": "dk2",
    "LIGHT_1": "lt1",
    "DARK_1": "dk1",
    "LIGHT_2": "lt2",
    "DARK_2": "dk2",
    **{f"ACCENT_{n}": f"accent{n}" for n in range(1, 7)},
}
OTHER_KEYWORDS = ("müll", "kicker", "feuer", "ww", "mcs")


def _theme_colors(presentation):
    theme = presentation.slide_masters[0].part.part_related_by(RT.THEME)
    root = etree.fromstring(theme.blob)
    colors = {}
    for slot in root.iter(qn("a:clrScheme")):
        for child in slot:
            color = child.find(qn("a:srgbClr"))
            value = color.get("val") if color is not None else child.find(qn("a:sysClr")).get("lastClr")
            colors[etree.QName(child).localname] = f"#{value.lower()}"
        break
    return colors


def _fill_color(shape, theme):
    try:
        if shape.fill.type != MSO_FILL.SOLID:
            return ""
        color = shape.fill.fore_color
        if color.type == MSO_COLOR_TYPE.RGB:
            value = f"#{str(color.rgb).lower()}"
        else:
            value = theme.get(THEME_SLOTS.get(color.theme_color.name, ""), "")
    except (AttributeError, KeyError, TypeError, ValueError):
        return ""
    # Schwarz gefüllte Räume der Skizze etwas aufhellen, damit die Schrift lesbar bleibt
    return "#495057" if value == "#000000" else value


def _iter_shapes(shapes, transform=(0, 0, 1, 1)):
    """Liefert (Form, Transformation). Gruppen haben eigene Kind-Koordinaten, die wir umrechnen."""
    a, b, sx, sy = transform
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            xfrm = shape._element.grpSpPr.find(qn("a:xfrm"))
            off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
            ch_off, ch_ext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
            gsx = int(ext.get("cx")) / (int(ch_ext.get("cx")) or 1)
            gsy = int(ext.get("cy")) / (int(ch_ext.get("cy")) or 1)
            inner = (
                a + (int(off.get("x")) - int(ch_off.get("x")) * gsx) * sx,
                b + (int(off.get("y")) - int(ch_off.get("y")) * gsy) * sy,
                sx * gsx,
                sy * gsy,
            )
            yield from _iter_shapes(shape.shapes, inner)
        else:
            yield shape, transform


def _guess_kind(label, color, width, height):
    text = label.lower()
    if "schrank" in text or "regal" in text:
        return Kind.CABINET
    if "tisch" in text:
        return Kind.WORKSTATION
    if not label and color == "#ffff00":
        return Kind.MARKING
    if width * height >= 10:
        return Kind.AREA
    if not label or any(word in text for word in OTHER_KEYWORDS):
        return Kind.OTHER
    return Kind.TEST_RIG


def read_elements(path, scale=1.0):
    """Liest die erste Folie. Gibt (Breite, Länge, Liste von Objekt-Dicts) in Metern zurück."""
    presentation = Presentation(path)
    theme = _theme_colors(presentation)
    to_m = scale / EMU_PER_CM
    slide_w, slide_h = presentation.slide_width * to_m, presentation.slide_height * to_m
    elements = []
    for z, (shape, (a, b, sx, sy)) in enumerate(_iter_shapes(presentation.slides[0].shapes)):
        if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE or shape.width is None:
            continue
        x, y = (a + shape.left * sx) * to_m, (b + shape.top * sy) * to_m
        w, h = shape.width * sx * to_m, shape.height * sy * to_m
        rotation = round(shape.rotation) % 360
        if rotation in (90, 270):
            # Achsenparallel ablegen: Mittelpunkt bleibt, Breite und Tiefe tauschen
            cx, cy = x + w / 2, y + h / 2
            w, h = h, w
            x, y = cx - w / 2, cy - h / 2
            rotation = 0
        elif rotation == 180:
            rotation = 0
        cx, cy = x + w / 2, y + h / 2
        if not (0 <= cx <= slide_w and 0 <= cy <= slide_h):
            continue
        label = shape.text_frame.text.replace("\v", "\n").strip() if shape.has_text_frame else ""
        color = _fill_color(shape, theme)
        elements.append(
            {
                "kind": _guess_kind(label, color, w, h),
                "label": label,
                "x": x,
                "y": y,
                "width": w,
                "height": h,
                "rotation": rotation,
                "color": color,
                "z": z,
            }
        )
    _number_duplicates(elements)
    return round(slide_w, 2), round(slide_h, 2), elements


def _number_duplicates(elements):
    """Gleich beschriftete Ablageorte durchnummerieren (von oben links nach unten rechts): „Schrank 1“ …"""
    groups = {}
    for element in elements:
        if element["kind"] in PlanElement.HOLDING_KINDS and element["label"]:
            groups.setdefault(element["label"], []).append(element)
    for label, group in groups.items():
        if len(group) > 1:
            group.sort(key=lambda e: (round(e["y"]), e["x"]))
            for number, element in enumerate(group, 1):
                element["label"] = f"{label} {number}"

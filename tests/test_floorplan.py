import json

import pytest
from django.core.management import call_command
from django.urls import reverse
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Cm

from floorplan.models import FloorPlan, PlanElement, element_for_location
from floorplan.pptx_import import read_elements
from floorplan.services import PlanError, save_plan
from inventory.models import Location

from .factories import ItemFactory, LocationFactory

Kind = PlanElement.Kind


def _add_rect(shapes, label, x, y, w, h, color=None, rotation=0):
    shape = shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    shape.text_frame.text = label
    shape.rotation = rotation
    if color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
    return shape


@pytest.fixture
def sketch(tmp_path):
    """Kleine Skizze: 10 × 20 cm (= m), mit Duplikaten, Drehung, Gruppe und Notizen neben der Folie."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(10), Cm(20)
    shapes = prs.slides.add_slide(prs.slide_layouts[6]).shapes
    _add_rect(shapes, "Werkstatt", 0, 0, 5, 4)
    _add_rect(shapes, "Schrank", 6, 5, 1, 0.5)
    _add_rect(shapes, "Schrank", 6, 1, 1, 0.5)
    _add_rect(shapes, "Labortisch", 2, 10, 2, 1, rotation=90)
    _add_rect(shapes, "", 8, 0, 0.1, 20, color="FFFF00")
    _add_rect(shapes, "Notiz", -5, 2, 3, 1)  # links neben der Folie
    shapes.add_textbox(Cm(1), Cm(15), Cm(3), Cm(1)).text_frame.text = "6700"
    group = shapes.add_group_shape()
    _add_rect(group.shapes, "Zwick", 3, 16, 1, 1, color="C00000")
    path = tmp_path / "skizze.pptx"
    prs.save(path)
    return path


def test_read_elements(sketch):
    width, height, elements = read_elements(sketch)
    assert (width, height) == (10, 20)
    by_label = {e["label"]: e for e in elements}
    assert set(by_label) == {"Werkstatt", "Schrank 1", "Schrank 2", "Labortisch", "", "Zwick"}
    # Duplikate von oben nach unten nummeriert
    assert by_label["Schrank 1"]["y"] == pytest.approx(1)
    assert by_label["Schrank 2"]["y"] == pytest.approx(5)
    # 90° gedreht: achsenparallel mit getauschten Maßen um denselben Mittelpunkt
    table = by_label["Labortisch"]
    assert (table["width"], table["height"]) == (pytest.approx(1), pytest.approx(2))
    assert (table["x"], table["y"]) == (pytest.approx(2.5), pytest.approx(9.5))
    assert table["rotation"] == 0
    assert by_label["Zwick"]["x"] == pytest.approx(3) and by_label["Zwick"]["color"] == "#c00000"
    assert by_label["Werkstatt"]["kind"] == Kind.AREA
    assert by_label["Schrank 1"]["kind"] == Kind.CABINET
    assert table["kind"] == Kind.WORKSTATION
    assert by_label["Zwick"]["kind"] == Kind.TEST_RIG
    assert by_label[""]["kind"] == Kind.MARKING


def test_import_command_creates_locations(db, sketch):
    call_command("import_floorplan", str(sketch), name="Halle")
    plan = FloorPlan.objects.get(name="Halle")
    assert plan.elements.count() == 6
    paths = set(Location.objects.filter(parent=plan.location).values_list("path", flat=True))
    assert paths == {
        "Halle › Werkstatt",
        "Halle › Schrank 1",
        "Halle › Schrank 2",
        "Halle › Labortisch",
        "Halle › Zwick",
    }
    # Ohne --replace kein zweiter Import
    with pytest.raises(Exception, match="replace"):
        call_command("import_floorplan", str(sketch), name="Halle")
    call_command("import_floorplan", str(sketch), name="Halle", replace=True)
    assert plan.elements.count() == 6


@pytest.fixture
def plan(db):
    return FloorPlan.objects.create(name="Halle", location=LocationFactory(name="Halle"), width=10, height=20)


def _element(**values):
    return {"kind": "cabinet", "label": "Schrank A", "x": 1, "y": 1, "width": 1, "height": 0.5, **values}


def test_save_plan_syncs_locations(plan):
    save_plan(plan, [_element()], [])
    element = plan.elements.get()
    assert element.location.path == "Halle › Schrank A"

    save_plan(plan, [_element(id=element.pk, label="Schrank B")], [])
    element.refresh_from_db()
    assert element.location.path == "Halle › Schrank B"

    # Markierung ist kein Ablageort: Ort wird entfernt
    location_pk = element.location_id
    save_plan(plan, [_element(id=element.pk, kind="marking")], [])
    element.refresh_from_db()
    assert element.location is None
    assert not Location.objects.filter(pk=location_pk).exists()


def test_same_label_gets_unique_location(plan):
    save_plan(plan, [_element(), _element(x=3)], [])
    names = sorted(plan.elements.values_list("location__name", flat=True))
    assert names == ["Schrank A", "Schrank A (2)"]


def test_cannot_delete_element_with_items(plan):
    save_plan(plan, [_element()], [])
    element = plan.elements.get()
    ItemFactory(location=element.location)
    with pytest.raises(PlanError, match="1 Gerät"):
        save_plan(plan, [], [element.pk])
    with pytest.raises(PlanError):
        save_plan(plan, [_element(id=element.pk, kind="marking")], [])
    assert plan.elements.filter(pk=element.pk).exists()

    # Leeres Objekt lässt sich löschen, sein Ort verschwindet mit
    save_plan(plan, [_element(label="Leer", x=5)], [])
    empty = plan.elements.get(label="Leer")
    save_plan(plan, [], [empty.pk])
    assert not Location.objects.filter(name="Leer").exists()


def test_element_for_sublocation(plan):
    save_plan(plan, [_element()], [])
    element = plan.elements.get()
    shelf = Location.objects.create(name="Fach 2", parent=element.location)
    assert element_for_location(shelf) == element


def test_pages_require_login(client, plan):
    response = client.get(reverse("floorplan:detail", args=[plan.pk]))
    assert response.status_code == 302 and "/konto/login/" in response["Location"]


def test_detail_highlights_item(client, user, plan):
    save_plan(plan, [_element()], [])
    element = plan.elements.get()
    item = ItemFactory(location=element.location)
    client.force_login(user)
    response = client.get(reverse("floorplan:detail", args=[plan.pk]), {"geraet": item.pk})
    assert response.status_code == 200
    data = response.context["plan_data"]
    assert data["highlight"] == element.pk
    assert data["elements"][0]["items"] == 1
    assert not data["editable"]


def test_panel_and_search(client, user, plan):
    save_plan(plan, [_element()], [])
    element = plan.elements.get()
    ItemFactory(name="Kraftsensor KMS", location=element.location)
    client.force_login(user)
    panel = client.get(reverse("floorplan:element", args=[element.pk]))
    assert "Kraftsensor KMS" in panel.content.decode()
    search = client.get(reverse("floorplan:search", args=[plan.pk]), {"q": "kraft"})
    assert f'data-show-element="{element.pk}"' in search.content.decode()


def test_only_staff_can_edit(client, user, staff, plan):
    client.force_login(user)
    assert client.get(reverse("floorplan:edit", args=[plan.pk])).status_code == 403
    payload = json.dumps({"elements": [_element()], "deleted": []})
    url = reverse("floorplan:save", args=[plan.pk])
    assert client.post(url, payload, content_type="application/json").status_code == 403

    client.force_login(staff)
    assert client.get(reverse("floorplan:edit", args=[plan.pk])).status_code == 200
    response = client.post(url, payload, content_type="application/json")
    assert response.status_code == 200
    assert response.json()["elements"][0]["label"] == "Schrank A"


def test_save_endpoint_reports_errors(client, staff, plan):
    client.force_login(staff)
    url = reverse("floorplan:save", args=[plan.pk])
    response = client.post(
        url, json.dumps({"elements": [{"kind": "cabinet"}]}), content_type="application/json"
    )
    assert response.status_code == 400
    assert "Ungültige" in response.json()["error"]
